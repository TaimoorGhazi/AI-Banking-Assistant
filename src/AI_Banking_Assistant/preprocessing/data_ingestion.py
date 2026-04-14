"""
Data ingestion module for AI Banking Assistant.
Loads documents from various file formats (txt, pdf, docx, json, jsonl, xlsx).
"""

import os
import json
import re
from pathlib import Path
from typing import Dict, List, Tuple

from src.AI_Banking_Assistant.core.logger import get_logger
from src.AI_Banking_Assistant.core.constants import SUPPORTED_FILE_EXTENSIONS, DATA_RAW_DIR
from src.AI_Banking_Assistant.core.exceptions import DataIngestionError

logger = get_logger(__name__)

QUESTION_HINTS = (
    "?",
    "what ",
    "how ",
    "can i ",
    "does ",
    "is ",
    "are ",
    "who ",
    "when ",
    "where ",
    "why ",
    "which ",
    "do i ",
)


def load_text_file(filepath: str) -> str:
    """Load content from a plain text file.
    
    Args:
        filepath: Path to the .txt file
        
    Returns:
        File content as string
        
    Raises:
        DataIngestionError: If file cannot be read
    """
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            content = f.read()
        logger.info(f"Loaded text file: {filepath} ({len(content)} chars)")
        return content
    except Exception as e:
        raise DataIngestionError(f"Failed to load text file {filepath}: {e}")


def load_pdf_file(filepath: str) -> str:
    """Load content from a PDF file.
    
    Args:
        filepath: Path to the .pdf file
        
    Returns:
        Extracted text content
        
    Raises:
        DataIngestionError: If file cannot be read
    """
    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(filepath)
        content = ""
        for page in reader.pages:
            text = page.extract_text()
            if text:
                content += text + "\n"
        logger.info(f"Loaded PDF: {filepath} ({len(reader.pages)} pages, {len(content)} chars)")
        return content
    except ImportError:
        raise DataIngestionError("PyPDF2 not installed. Run: pip install PyPDF2")
    except Exception as e:
        raise DataIngestionError(f"Failed to load PDF {filepath}: {e}")


def load_docx_file(filepath: str) -> str:
    """Load content from a Word document.
    
    Args:
        filepath: Path to the .docx file
        
    Returns:
        Extracted text content
        
    Raises:
        DataIngestionError: If file cannot be read
    """
    try:
        from docx import Document

        doc = Document(filepath)
        content = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        logger.info(f"Loaded DOCX: {filepath} ({len(doc.paragraphs)} paragraphs, {len(content)} chars)")
        return content
    except ImportError:
        raise DataIngestionError("python-docx not installed. Run: pip install python-docx")
    except Exception as e:
        raise DataIngestionError(f"Failed to load DOCX {filepath}: {e}")


def _flatten_json_to_text(data: object) -> str:
    """Convert structured JSON content into retrieval-friendly plain text."""
    lines: List[str] = []

    def walk(node: object):
        if isinstance(node, dict):
            # Handle common FAQ shape: {"question": ..., "answer": ...}
            if "question" in node and "answer" in node:
                question = str(node.get("question", "")).strip()
                answer = str(node.get("answer", "")).strip()
                if question or answer:
                    lines.append(f"Q: {question}\nA: {answer}")

            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for item in node:
                walk(item)
        elif isinstance(node, str):
            text = node.strip()
            if text:
                lines.append(text)

    walk(data)
    # Deduplicate while preserving order.
    seen = set()
    deduped = []
    for line in lines:
        if line not in seen:
            deduped.append(line)
            seen.add(line)
    return "\n\n".join(deduped)


def load_json_file(filepath: str) -> str:
    """Load and flatten content from a JSON file."""
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            data = json.load(f)
        content = _flatten_json_to_text(data)
        logger.info(f"Loaded JSON: {filepath} ({len(content)} chars)")
        return content
    except Exception as e:
        raise DataIngestionError(f"Failed to load JSON {filepath}: {e}")


def load_jsonl_file(filepath: str) -> str:
    """Load and flatten content from a JSONL file."""
    try:
        entries: List[object] = []
        with open(filepath, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                entries.append(json.loads(line))

        content = _flatten_json_to_text(entries)
        logger.info(f"Loaded JSONL: {filepath} ({len(entries)} rows, {len(content)} chars)")
        return content
    except Exception as e:
        raise DataIngestionError(f"Failed to load JSONL {filepath}: {e}")


def _normalize_excel_text(value: object) -> str:
    """Normalize cell text extracted from an Excel worksheet."""
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\u200b", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def _looks_like_excel_question(text: str) -> bool:
    """Heuristic check for identifying question-style rows in workbook sheets."""
    lowered = text.lower().strip()
    return bool(lowered) and (
        lowered.endswith("?") or any(lowered.startswith(prefix) for prefix in QUESTION_HINTS)
    )


def _parse_excel_sheet(worksheet) -> Tuple[List[Tuple[str, str]], List[str]]:
    """Extract Q/A pairs and useful tabular lines from one worksheet."""
    rows: List[List[str]] = []
    for row in worksheet.iter_rows(values_only=True):
        normalized_cells = [_normalize_excel_text(cell) for cell in row]
        cells = [cell for cell in normalized_cells if cell]
        if cells:
            rows.append(cells)

    pairs: List[Tuple[str, str]] = []
    tabular_lines: List[str] = []
    consumed_rows = set()

    current_question = ""
    current_answer_parts: List[str] = []

    for idx, cells in enumerate(rows):
        first_cell = cells[0]
        second_cell = cells[1] if len(cells) > 1 else ""

        if _looks_like_excel_question(first_cell):
            if current_question and current_answer_parts:
                answer = " ".join(current_answer_parts).strip()
                if answer:
                    pairs.append((current_question, answer))

            current_question = first_cell
            current_answer_parts = []
            consumed_rows.add(idx)

            if second_cell and second_cell.lower() != "main" and not _looks_like_excel_question(second_cell):
                current_answer_parts.append(second_cell)

            for extra in cells[2:]:
                if extra and not _looks_like_excel_question(extra):
                    current_answer_parts.append(extra)
            continue

        if current_question:
            current_answer_parts.extend(cells)
            consumed_rows.add(idx)

    if current_question and current_answer_parts:
        answer = " ".join(current_answer_parts).strip()
        if answer:
            pairs.append((current_question, answer))

    for idx, cells in enumerate(rows):
        if idx in consumed_rows:
            continue
        line = " | ".join(cells).strip()
        if line:
            tabular_lines.append(line)

    dedup_pairs: List[Tuple[str, str]] = []
    seen_pairs = set()
    for question, answer in pairs:
        key = (question, answer)
        if key in seen_pairs:
            continue
        dedup_pairs.append((question, answer))
        seen_pairs.add(key)

    dedup_lines: List[str] = []
    seen_lines = set()
    for line in tabular_lines:
        if line in seen_lines:
            continue
        dedup_lines.append(line)
        seen_lines.add(line)

    return dedup_pairs, dedup_lines


def _extract_main_sheet_hyperlinks(workbook) -> List[str]:
    """Extract main-sheet product hyperlink mapping to capture workbook navigation context."""
    main_sheet = None
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().lower() == "main":
            main_sheet = workbook[sheet_name]
            break

    if main_sheet is None:
        return []

    lines: List[str] = []
    for row in main_sheet.iter_rows(min_row=1, max_row=500):
        for cell in row:
            if cell.hyperlink is None:
                continue

            label = _normalize_excel_text(cell.value)
            target = cell.hyperlink.target or cell.hyperlink.location or ""
            if "!" in target:
                target = target.split("!", 1)[0]
            target = target.strip("'").strip()

            if label:
                if target:
                    lines.append(f"Main index maps '{label}' to sheet '{target}'.")
                else:
                    lines.append(f"Main index product: {label}")

    deduped_lines: List[str] = []
    seen = set()
    for line in lines:
        if line in seen:
            continue
        deduped_lines.append(line)
        seen.add(line)
    return deduped_lines


def load_xlsx_file(filepath: str) -> str:
    """Load and flatten content from an Excel workbook for retrieval."""
    try:
        from openpyxl import load_workbook

        # Keep read_only disabled so hyperlink targets are available on cells.
        workbook = load_workbook(filepath, read_only=False, data_only=True)
        sections: List[str] = []

        navigation_lines = _extract_main_sheet_hyperlinks(workbook)
        if navigation_lines:
            sections.append("[Workbook Navigation]\n" + "\n".join(navigation_lines))

        for worksheet in workbook.worksheets:
            sheet_title = _normalize_excel_text(worksheet.title) or "Unnamed Sheet"
            qa_pairs, tabular_lines = _parse_excel_sheet(worksheet)

            sheet_lines: List[str] = [f"[Sheet: {sheet_title}]"]
            for question, answer in qa_pairs:
                sheet_lines.append(f"Q: {question}\nA: {answer}")

            # Keep non-QA rows too so pricing/rate-table sheets remain searchable.
            if tabular_lines:
                sheet_lines.extend(tabular_lines)

            sections.append("\n\n".join(sheet_lines))

        content = "\n\n---\n\n".join(section for section in sections if section.strip())
        logger.info(
            "Loaded XLSX: %s (%d sheets, %d chars)",
            filepath,
            len(workbook.sheetnames),
            len(content),
        )
        return content
    except ImportError:
        raise DataIngestionError("openpyxl not installed. Run: pip install openpyxl")
    except Exception as e:
        raise DataIngestionError(f"Failed to load XLSX {filepath}: {e}")


def load_document(filepath: str) -> str:
    """Load a document based on its file extension.
    
    Args:
        filepath: Path to the document
        
    Returns:
        Document content as string
        
    Raises:
        DataIngestionError: If format not supported or file cannot be read
    """
    ext = Path(filepath).suffix.lower()

    if ext == ".txt":
        return load_text_file(filepath)
    elif ext == ".pdf":
        return load_pdf_file(filepath)
    elif ext == ".docx":
        return load_docx_file(filepath)
    elif ext == ".json":
        return load_json_file(filepath)
    elif ext == ".jsonl":
        return load_jsonl_file(filepath)
    elif ext == ".xlsx":
        return load_xlsx_file(filepath)
    else:
        raise DataIngestionError(
            f"Unsupported file format: {ext}. Supported: {SUPPORTED_FILE_EXTENSIONS}"
        )


def ingest_directory(dir_path: str = None) -> List[Dict[str, str]]:
    """Load all supported documents from a directory.
    
    Args:
        dir_path: Path to directory (defaults to data/raw/)
        
    Returns:
        List of dicts with 'filename', 'filepath', and 'content' keys
        
    Raises:
        DataIngestionError: If directory does not exist
    """
    if dir_path is None:
        dir_path = DATA_RAW_DIR

    dir_path = Path(dir_path)
    if not dir_path.exists():
        raise DataIngestionError(f"Directory does not exist: {dir_path}")

    documents = []
    supported_exts = set(SUPPORTED_FILE_EXTENSIONS)

    for filepath in sorted(dir_path.iterdir()):
        if filepath.is_file() and filepath.suffix.lower() in supported_exts:
            try:
                content = load_document(str(filepath))
                documents.append({
                    "filename": filepath.name,
                    "filepath": str(filepath),
                    "content": content,
                })
            except DataIngestionError as e:
                logger.warning(f"Skipping {filepath.name}: {e}")

    logger.info(f"Ingested {len(documents)} documents from {dir_path}")
    return documents


if __name__ == "__main__":
    # Quick test
    docs = ingest_directory()
    for doc in docs:
        print(f"  {doc['filename']}: {len(doc['content'])} chars")
