"""Convert the NUST Bank knowledge workbook into instruction JSONL files."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Iterable, List, Optional, Tuple

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.AI_Banking_Assistant.llm.finetuning.prepare_dataset import save_jsonl

QUESTION_HINTS = (
    "?",
    "what ",
    "how ",
    "can i ",
    "does ",
    "is ",
    "are ",
    "who ",
    "when ",
    "where ",
    "why ",
    "which ",
    "do i ",
)

SKIP_SHEET_NAMES = {"main", "rate sheet july 1 2024", "sheet1"}


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\u200b", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def looks_like_question(text: str) -> bool:
    lowered = text.lower().strip()
    return bool(lowered) and (lowered.endswith("?") or any(lowered.startswith(prefix) for prefix in QUESTION_HINTS))


def extract_pairs_from_sheet(worksheet) -> List[Tuple[str, str, str]]:
    pairs: List[Tuple[str, str, str]] = []
    rows = list(worksheet.iter_rows(values_only=True))
    sheet_name = worksheet.title.strip()

    current_question: Optional[str] = None
    current_answer_parts: List[str] = []

    for row in rows:
        cells = [normalize_text(cell) for cell in row if normalize_text(cell)]
        if not cells:
            continue

        # Skip obvious navigation or title-only rows.
        if len(cells) == 1 and cells[0].lower() in {"main"}:
            continue

        first_cell = cells[0]
        second_cell = cells[1] if len(cells) > 1 else ""

        if looks_like_question(first_cell):
            if current_question and current_answer_parts:
                pairs.append((sheet_name, current_question, " ".join(current_answer_parts).strip()))
            current_question = first_cell
            current_answer_parts = []
            if second_cell and not looks_like_question(second_cell):
                current_answer_parts.append(second_cell)
            for extra in cells[2:]:
                if extra and not looks_like_question(extra):
                    current_answer_parts.append(extra)
            continue

        if current_question is None:
            continue

        if len(cells) == 1:
            current_answer_parts.append(cells[0])
        else:
            # Treat non-question rows as answer continuation.
            current_answer_parts.extend(cells)

    if current_question and current_answer_parts:
        pairs.append((sheet_name, current_question, " ".join(current_answer_parts).strip()))

    # Remove duplicates while preserving order.
    seen = set()
    deduped: List[Tuple[str, str, str]] = []
    for item in pairs:
        key = (item[1], item[2])
        if key in seen:
            continue
        seen.add(key)
        deduped.append(item)
    return deduped


def build_samples(workbook_path: Path) -> List[dict]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    samples: List[dict] = []

    for worksheet in workbook.worksheets:
        if worksheet.title.strip().lower() in SKIP_SHEET_NAMES:
            continue
        for category, question, answer in extract_pairs_from_sheet(worksheet):
            if not question or not answer:
                continue
            samples.append(
                {
                    "instruction": question,
                    "context": f"Category: {category}",
                    "response": answer,
                    "source": worksheet.title,
                }
            )

    return samples


def split_samples(samples: List[dict], train_ratio: float = 0.7, val_ratio: float = 0.15) -> Tuple[List[dict], List[dict], List[dict]]:
    total = len(samples)
    train_end = int(total * train_ratio)
    val_end = train_end + int(total * val_ratio)
    return samples[:train_end], samples[train_end:val_end], samples[val_end:]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert a workbook into JSONL training data.")
    parser.add_argument("--input", default=str(PROJECT_ROOT / "NUST Bank-Product-Knowledge (1).xlsx"), help="Path to the xlsx workbook")
    parser.add_argument("--output-dir", default=str(PROJECT_ROOT / "data" / "finetune"), help="Output directory for train/val/test JSONL files")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.input)
    output_dir = Path(args.output_dir)

    samples = build_samples(workbook_path)
    if not samples:
        raise RuntimeError(f"No question/answer pairs were extracted from {workbook_path}")

    train_samples, val_samples, test_samples = split_samples(samples)
    output_dir.mkdir(parents=True, exist_ok=True)
    save_jsonl(train_samples, str(output_dir / "train.jsonl"))
    save_jsonl(val_samples, str(output_dir / "val.jsonl"))
    save_jsonl(test_samples, str(output_dir / "test.jsonl"))

    print(json.dumps({
        "input": str(workbook_path),
        "output_dir": str(output_dir),
        "total": len(samples),
        "train": len(train_samples),
        "val": len(val_samples),
        "test": len(test_samples),
    }, indent=2))


if __name__ == "__main__":
    main()
